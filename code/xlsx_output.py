# -*- coding: utf-8 -*-
"""Excel result-table writer used by the four complete programs."""
from pathlib import Path
from openpyxl import Workbook

def write_table(path: Path, sheets):
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    book.remove(book.active)
    for title, header, rows in sheets:
        ws = book.create_sheet(title)
        ws.append(list(header))
        for row in rows:
            ws.append(list(row))
        ws.freeze_panes = 'B2'
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = cell.font.copy(bold=True)
        ws.column_dimensions['A'].width = 18
        for col in range(2, ws.max_column + 1):
            ws.column_dimensions[ws.cell(1, col).column_letter].width = 12
    book.save(path)

